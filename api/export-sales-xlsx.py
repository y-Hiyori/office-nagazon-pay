#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Vercel Python Function: /api/export-sales-xlsx
BaseHTTPRequestHandler 版。Vercel の /api ディレクトリ用 Python 関数として公開する。

シート構成（見やすい順）:
  1. 要点まとめ   … 開いてすぐ分かる要点ダッシュボード
  2. 売上サマリー … 入金売上・粗利などのくわしい内訳
  3. 期間比較     … 今期 vs 比較期間
  4. 商品別売上   … 商品ごとの売上・原価・粗利 ＋ グラフ
  5. 注文一覧     … 注文ごとの明細
  6. 処分履歴     … 廃棄ロスの明細（該当があるときだけ）

デザイン: ネイビー×スレートの落ち着いた配色。見出しは淡色＋ネイビー文字、
強調はアクセントブルー、黒字＝緑／赤字＝赤で統一。
"""
import io
import json
from http.server import BaseHTTPRequestHandler
from urllib.parse import quote
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

# ---- 配色（プロ向けパレット） ----------------------------------------------
NAVY = "0F172A"        # 濃紺（文字・タイトル帯）
SLATE_700 = "334155"   # 本文
SLATE_500 = "64748B"   # 補助文字
SLATE_300 = "CBD5E1"   # 中罫線
SLATE_200 = "E2E8F0"   # 細罫線
SLATE_100 = "F1F5F9"   # 見出し・合計の淡色
SLATE_50 = "F8FAFC"    # ゼブラ
WHITE = "FFFFFF"
ACCENT = "1D4ED8"      # アクセントブルー
ACCENT_SOFT = "DBEAFE"  # アクセント淡色
POS = "047857"         # 黒字（緑）
NEG = "BE123C"         # 赤字（赤）
POS_BG = "ECFDF5"      # 強調カードの淡緑
NEG_BG = "FFF1F2"      # 注意の淡赤

TITLE_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FILL = PatternFill("solid", fgColor=SLATE_100)
ZEBRA_FILL = PatternFill("solid", fgColor=SLATE_50)
TOTAL_FILL = PatternFill("solid", fgColor=SLATE_100)
HIGHLIGHT_FILL = PatternFill("solid", fgColor=POS_BG)
SECTION_FILL = PatternFill("solid", fgColor=SLATE_100)
KPI_FILL = PatternFill("solid", fgColor=SLATE_50)

TITLE_FONT = Font(name="メイリオ", size=15, bold=True, color=WHITE)
HEADER_FONT = Font(name="メイリオ", size=10.5, bold=True, color=NAVY)
BODY_FONT = Font(name="メイリオ", size=10, color=SLATE_700)
BOLD_FONT = Font(name="メイリオ", size=10, bold=True, color=NAVY)
NOTE_FONT = Font(name="メイリオ", size=9, color=SLATE_500)
SECTION_FONT = Font(name="メイリオ", size=11, bold=True, color=ACCENT)
KPI_LABEL_FONT = Font(name="メイリオ", size=10, bold=True, color=SLATE_500)
KPI_VALUE_FONT = Font(name="メイリオ", size=20, bold=True, color=NAVY)
KPI_SUB_FONT = Font(name="メイリオ", size=14, bold=True, color=NAVY)
GOOD_FONT = Font(name="メイリオ", size=10, bold=True, color=POS)
BAD_FONT = Font(name="メイリオ", size=10, bold=True, color=NEG)

THIN = Side(style="thin", color=SLATE_200)
MED = Side(style="medium", color=SLATE_300)
ACCENT_SIDE = Side(style="thick", color=ACCENT)
BORDER = Border(top=THIN, left=THIN, bottom=THIN, right=THIN)
HEADER_BORDER = Border(top=THIN, left=THIN, bottom=MED, right=THIN)
TOTAL_BORDER = Border(top=MED, left=THIN, bottom=THIN, right=THIN)
CARD_BORDER = Border(top=MED, left=THIN, bottom=THIN, right=THIN)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER_NO_WRAP = Alignment(horizontal="center", vertical="center")

YEN = "¥#,##0"
COUNT = "#,##0"
PCT = "0.0%"


def to_int(v):
    try:
        return int(round(float(v or 0)))
    except Exception:
        return 0


def set_number_format(cell, fmt):
    if fmt == "yen":
        cell.number_format = YEN
    elif fmt == "count":
        cell.number_format = COUNT
    elif fmt == "pct":
        cell.number_format = PCT


def write_headers(ws, row, headers, height=None):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = HEADER_BORDER
    if height:
        ws.row_dimensions[row].height = height


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_title_bar(ws, row, text, span, height=34):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL
    cell.alignment = LEFT
    ws.row_dimensions[row].height = height


def draw_section(ws, row, text, span):
    """淡色ベタ＋左にアクセントバーを引いた見出し行。"""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    for c in range(1, span + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = SECTION_FILL
        cell.border = Border(top=THIN, bottom=THIN,
                             left=(ACCENT_SIDE if c == 1 else None),
                             right=(THIN if c == span else None))
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24


def write_notes(ws, start_row, span, notes):
    for i, txt in enumerate(notes):
        r = start_row + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
        cell = ws.cell(row=r, column=1, value=txt)
        cell.font = NOTE_FONT
        cell.alignment = LEFT
        ws.row_dimensions[r].height = 18


def style_range(ws, start_row, end_row, formats=None, center_cols=()):
    formats = formats or {}
    center_set = set(center_cols)
    for r in range(start_row, end_row + 1):
        zebra = ((r - start_row) % 2) == 1
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            if zebra:
                cell.fill = ZEBRA_FILL
            cell.border = BORDER
            if c in center_set:
                cell.alignment = CENTER
            elif formats.get(c) in ("yen", "count", "pct"):
                cell.alignment = RIGHT
            else:
                cell.alignment = LEFT
            set_number_format(cell, formats.get(c))


def total_row_style(cell, align=RIGHT, fmt=None):
    cell.font = BOLD_FONT
    cell.fill = TOTAL_FILL
    cell.alignment = align
    cell.border = TOTAL_BORDER
    set_number_format(cell, fmt)


# ---------------------------------------------------------------------------
# 1. 要点まとめ（先頭シート）
# ---------------------------------------------------------------------------
def build_overview_sheet(wb, range_label, now_jst, summary, products, prev):
    ws = wb.create_sheet("要点まとめ")
    ncol = 6
    set_widths(ws, [26, 18, 20, 18, 20, 18])

    write_title_bar(ws, 1, "OFFICE NAGAZON 売上レポート｜要点まとめ", ncol, height=36)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    sub = ws.cell(row=2, column=1,
                  value="対象期間：{}　　出力日時：{}".format(range_label, now_jst))
    sub.font = NOTE_FONT
    sub.alignment = LEFT
    ws.row_dimensions[2].height = 20

    cash = to_int(summary.get("cashSales", 0))
    profit = to_int(summary.get("profitTotal", 0))
    cost = to_int(summary.get("costTotal", 0))
    disposal = to_int(summary.get("disposalCost", 0))
    gross = to_int(summary.get("grossSubtotal", 0))
    coupon = to_int(summary.get("couponDiscount", 0))
    points = to_int(summary.get("pointsTotal", 0))
    orders_n = to_int(summary.get("orderCount", 0))
    qty = sum(to_int(p.get("quantity", 0)) for p in products)
    unit_price = int(round(cash / orders_n)) if orders_n else 0
    profit_rate = (profit / cash) if cash else 0
    profit_font = Font(name="メイリオ", size=20, bold=True, color=POS if profit >= 0 else NEG)

    def card(label_row, col, label, value, fmt, fill, value_font=KPI_VALUE_FONT):
        ws.merge_cells(start_row=label_row, start_column=col, end_row=label_row, end_column=col + 1)
        ws.merge_cells(start_row=label_row + 1, start_column=col, end_row=label_row + 1, end_column=col + 1)
        lc = ws.cell(row=label_row, column=col, value=label)
        lc.font = KPI_LABEL_FONT
        lc.alignment = CENTER_NO_WRAP
        vc = ws.cell(row=label_row + 1, column=col, value=value)
        vc.font = value_font
        vc.alignment = CENTER_NO_WRAP
        set_number_format(vc, fmt)
        for rr in (label_row, label_row + 1):
            for cc in (col, col + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.fill = fill
                cell.border = Border(
                    top=(CARD_BORDER.top if rr == label_row else None),
                    bottom=(CARD_BORDER.bottom if rr == label_row + 1 else None),
                    left=(CARD_BORDER.left if cc == col else None),
                    right=(CARD_BORDER.right if cc == col + 1 else None),
                )

    def mini_card(label_row, col, label, value, fmt):
        ws.merge_cells(start_row=label_row, start_column=col, end_row=label_row, end_column=col + 1)
        ws.merge_cells(start_row=label_row + 1, start_column=col, end_row=label_row + 1, end_column=col + 1)
        lc = ws.cell(row=label_row, column=col, value=label)
        lc.font = KPI_LABEL_FONT
        lc.alignment = CENTER_NO_WRAP
        vc = ws.cell(row=label_row + 1, column=col, value=value)
        vc.font = KPI_SUB_FONT
        vc.alignment = CENTER_NO_WRAP
        set_number_format(vc, fmt)
        for rr in (label_row, label_row + 1):
            for cc in (col, col + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.fill = KPI_FILL
                cell.border = Border(
                    top=(CARD_BORDER.top if rr == label_row else None),
                    bottom=(CARD_BORDER.bottom if rr == label_row + 1 else None),
                    left=(CARD_BORDER.left if cc == col else None),
                    right=(CARD_BORDER.right if cc == col + 1 else None),
                )

    r = 4
    draw_section(ws, r, "◆ いちばん大事な数字", ncol); r += 1
    card(r, 1, "入金売上（実際に入ったお金）", cash, "yen", HIGHLIGHT_FILL)
    card(r, 3, "粗利（利益）", profit, "yen", HIGHLIGHT_FILL, value_font=profit_font)
    card(r, 5, "粗利率", profit_rate, "pct", KPI_FILL)
    ws.row_dimensions[r].height = 20
    ws.row_dimensions[r + 1].height = 32
    r += 2

    draw_section(ws, r, "◆ お客様の動き", ncol); r += 1
    mini_card(r, 1, "注文件数", orders_n, "count")
    mini_card(r, 3, "販売個数", qty, "count")
    mini_card(r, 5, "平均注文額（客単価）", unit_price, "yen")
    ws.row_dimensions[r].height = 20
    ws.row_dimensions[r + 1].height = 26
    r += 2

    if prev:
        prev_label = prev.get("label") or "比較期間"
        draw_section(ws, r, "◆ 前の期間との比較（{} とくらべる）".format(prev_label), ncol); r += 1
        write_headers(ws, r, ["指標", "今期", prev_label, "増減率"], height=24)
        r += 1
        cur = prev.get("current") or {}
        pre = prev.get("previous") or {}
        for lab, cv, pv, fmt in [
            ("入金売上", to_int(cur.get("cashSales", 0)), to_int(pre.get("cashSales", 0)), "yen"),
            ("注文件数", to_int(cur.get("orderCount", 0)), to_int(pre.get("orderCount", 0)), "count"),
            ("商品売上（割引前）", to_int(cur.get("grossSubtotal", 0)), to_int(pre.get("grossSubtotal", 0)), "yen"),
        ]:
            ws.cell(row=r, column=1, value=lab).font = BOLD_FONT
            c2 = ws.cell(row=r, column=2, value=cv); set_number_format(c2, fmt); c2.alignment = RIGHT; c2.font = BODY_FONT
            c3 = ws.cell(row=r, column=3, value=pv); set_number_format(c3, fmt); c3.alignment = RIGHT; c3.font = BODY_FONT
            rate = ((cv - pv) / pv) if pv else 0
            c4 = ws.cell(row=r, column=4, value=rate)
            c4.number_format = PCT
            c4.alignment = RIGHT
            c4.font = GOOD_FONT if rate >= 0 else BAD_FONT
            for cc in range(1, 5):
                ws.cell(row=r, column=cc).border = BORDER
            ws.row_dimensions[r].height = 22
            r += 1
        r += 1

    draw_section(ws, r, "◆ 売上のお金の流れ", ncol); r += 1
    flow = [
        ("商品売上（割引前）", gross, False),
        ("クーポン割引（引く）", -coupon, False),
        ("ポイント充当（引く）", -points, False),
        ("＝ 入金売上", cash, True),
        ("売上原価（仕入れ・引く）", -cost, False),
        ("廃棄ロス（引く）", -disposal, False),
        ("＝ 粗利（利益）", profit, True),
    ]
    for lab, val, strong in flow:
        lc = ws.cell(row=r, column=1, value=lab)
        lc.font = BOLD_FONT if strong else BODY_FONT
        lc.alignment = LEFT
        lc.border = Border(top=THIN, bottom=THIN, left=(ACCENT_SIDE if strong else THIN), right=THIN)
        vc = ws.cell(row=r, column=2, value=val)
        vc.number_format = YEN
        vc.alignment = RIGHT
        vc.font = BOLD_FONT
        vc.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        if strong:
            lc.fill = HIGHLIGHT_FILL
            vc.fill = HIGHLIGHT_FILL
        ws.row_dimensions[r].height = 22
        r += 1

    r += 1
    write_notes(ws, r, ncol, [
        "「入金売上」はお客様から実際にいただいた金額です。",
        "「粗利（利益）」＝ 入金売上 − 売上原価（仕入れ） − 廃棄ロス です。",
        "くわしい内訳は「売上サマリー」、商品ごとは「商品別売上」、注文ごとは「注文一覧」シートを見てください。",
    ])
    ws.freeze_panes = "A3"
    return ws


# ---------------------------------------------------------------------------
# 2. 売上サマリー（くわしい内訳）
# ---------------------------------------------------------------------------
def build_summary_sheet(wb, range_label, now_jst, summary):
    ws = wb.create_sheet("売上サマリー")
    set_widths(ws, [28, 18, 62])
    write_title_bar(ws, 1, "売上サマリー（くわしい内訳）", 3, height=34)
    write_headers(ws, 2, ["項目", "値", "説明"], height=26)

    rows = [
        ("対象期間", range_label, "集計対象の期間"),
        ("出力日時", now_jst, "Excelを出力した日時"),
        ("商品売上（割引前）", to_int(summary.get("grossSubtotal", 0)), "商品ごとの売上計算の合計"),
        ("クーポン割引", to_int(summary.get("couponDiscount", 0)), "{0} 件で使用（引く）".format(to_int(summary.get("couponOrderCount", 0)))),
        ("ポイント充当", to_int(summary.get("pointsTotal", 0)), "{0} 件で使用（引く）".format(to_int(summary.get("pointsOrderCount", 0)))),
        ("入金売上", None, "＝ 商品売上 − クーポン割引 − ポイント充当"),
        ("総値引額", None, "＝ クーポン割引 ＋ ポイント充当"),
        ("実収率", None, "＝ 入金売上 ÷ 商品売上（割引前）"),
        ("売上原価（仕入れ）", to_int(summary.get("costTotal", 0)), "＝ Σ（仕入れ原価 × 販売数量）"),
        ("廃棄ロス（処分原価）", to_int(summary.get("disposalCost", 0)),
         "＝ Σ（処分した原価）／処分 {0} 個（引く）".format(to_int(summary.get("disposalQty", 0)))),
        ("粗利（利益）", None, "＝ 入金売上 − 売上原価（仕入れ） − 廃棄ロス"),
        ("粗利率", None, "＝ 粗利 ÷ 入金売上"),
    ]

    for i, (label, val, note) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label)
        if val is not None:
            ws.cell(row=i, column=2, value=val)
        ws.cell(row=i, column=3, value=note)

    style_range(ws, 3, 14, {2: "yen"})
    ws["B8"] = "=B5-B6-B7"
    ws["B9"] = "=B6+B7"
    ws["B10"] = "=IF(B5=0,0,B8/B5)"
    ws["B13"] = "=B8-B11-B12"
    ws["B14"] = "=IF(B8=0,0,B13/B8)"
    for addr, fmt in [("B8", YEN), ("B9", YEN), ("B10", PCT), ("B11", YEN), ("B12", YEN), ("B13", YEN), ("B14", PCT)]:
        ws[addr].number_format = fmt
        ws[addr].alignment = RIGHT
        ws[addr].font = Font(name="メイリオ", size=11 if addr in ("B8", "B13") else 10, bold=True,
                             color=NAVY)
    ws["B8"].fill = HIGHLIGHT_FILL
    ws["B13"].fill = HIGHLIGHT_FILL
    for addr in ("B8", "B13"):
        ws[addr].border = Border(top=MED, bottom=THIN, left=THIN, right=THIN)

    notes = [
        "入金売上 ＝ お客様から実際にいただいた金額です。",
        "粗利（利益）＝ 入金売上 − 売上原価（仕入れ） − 廃棄ロス（処分した原価）です。仕入れ原価は商品ごと・入荷ロットごとに管理画面で設定できます。",
        "廃棄ロスは「処分履歴」シートで理由と商品名を確認できます。",
        "商品別売上シートには売上・粗利・数量・売上構成比のグラフを追加しています。",
        "期間比較シートには今期 vs 比較期間の棒グラフと、増減率の折れ線グラフを追加しています。",
    ]
    write_notes(ws, 16, 3, notes)
    ws.freeze_panes = "A3"
    return ws


# ---------------------------------------------------------------------------
# 3. 期間比較
# ---------------------------------------------------------------------------
def build_compare_sheet(wb, range_label, prev):
    ws = wb.create_sheet("期間比較")
    prev_label = (prev or {}).get("label") or "比較期間"
    set_widths(ws, [28, 16, 16, 16, 14, 3, 14, 14, 14, 14, 14, 14, 14])
    write_title_bar(ws, 1, "期間比較（{0}{1}）".format(range_label, (" vs " + prev_label) if prev else ""), 5, height=36)
    write_headers(ws, 2, ["指標", "今期", prev_label, "増減", "増減率"], height=30)

    cur = (prev or {}).get("current") or {}
    pre = (prev or {}).get("previous") or {}
    rows = [
        ("入金売上", to_int(cur.get("cashSales", 0)), to_int(pre.get("cashSales", 0)), "yen"),
        ("注文件数", to_int(cur.get("orderCount", 0)), to_int(pre.get("orderCount", 0)), "count"),
        ("商品売上（割引前）", to_int(cur.get("grossSubtotal", 0)), to_int(pre.get("grossSubtotal", 0)), "yen"),
    ]
    for i, (label, curv, prevv, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = BOLD_FONT
        ws.cell(row=i, column=2, value=curv)
        ws.cell(row=i, column=3, value=prevv)
        ws.cell(row=i, column=4, value=f"=B{i}-C{i}")
        ws.cell(row=i, column=5, value=f"=IF(C{i}=0,0,D{i}/C{i})")
        style_range(ws, i, i, {2: fmt, 3: fmt, 4: fmt, 5: "pct"})
        ws.row_dimensions[i].height = 24
        diff = curv - prevv
        diff_font = GOOD_FONT if diff >= 0 else BAD_FONT
        ws.cell(row=i, column=4).font = diff_font
        ws.cell(row=i, column=5).font = diff_font

    if prev:
        compare_chart = BarChart()
        compare_chart.type = "col"
        compare_chart.style = 10
        compare_chart.title = "今期 vs {}".format(prev_label)
        compare_chart.legend.position = "b"
        compare_chart.gapWidth = 50
        compare_chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=2, max_row=5), titles_from_data=True)
        compare_chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=5))
        compare_chart.width = 13.8
        compare_chart.height = 8.2
        if len(compare_chart.series) >= 2:
            compare_chart.series[0].graphicalProperties.solidFill = ACCENT
            compare_chart.series[1].graphicalProperties.solidFill = SLATE_300
        ws.add_chart(compare_chart, "G3")

        rate_chart = LineChart()
        rate_chart.title = "増減率"
        rate_chart.style = 12
        rate_chart.legend = None
        rate_chart.add_data(Reference(ws, min_col=5, min_row=2, max_row=5), titles_from_data=True)
        rate_chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=5))
        rate_chart.width = 13.8
        rate_chart.height = 7.4
        if rate_chart.series:
            rate_chart.series[0].graphicalProperties.line.solidFill = NEG
            rate_chart.series[0].graphicalProperties.line.width = 24000
        ws.add_chart(rate_chart, "G20")

    ws.merge_cells("A8:E8")
    ws["A8"] = f"増減 ＝ 今期 − {prev_label}、増減率 ＝ 増減 ÷ {prev_label} です。緑＝増加、赤＝減少。"
    ws["A8"].font = NOTE_FONT
    ws["A8"].alignment = LEFT
    ws.row_dimensions[8].height = 22
    ws.auto_filter.ref = "A2:E5"
    ws.freeze_panes = "A3"
    return ws


# ---------------------------------------------------------------------------
# 4. 商品別売上（表＋ネイティブグラフ）
# ---------------------------------------------------------------------------
def build_product_sheet(wb, products):
    ws = wb.create_sheet("商品別売上")
    DATA_COLS = 13
    set_widths(ws, [34, 13, 8, 16, 15, 15, 17, 15, 13, 14, 11, 10, 11, 3, 3] + [14] * 6)
    write_title_bar(ws, 1, "商品別売上（売上・仕入れ原価・粗利がひと目で分かる一覧）", DATA_COLS, height=38)
    write_headers(ws, 2, [
        "商品名", "売価(平均・円)", "数量", "売上計算(割引前・円)", "クーポン割引(円)",
        "ポイント充当(円)", "入金売上(割引後・円)", "仕入れ原価(単価・円)", "売上原価(円)",
        "1個あたり粗利(円)", "粗利(円)", "粗利率", "売上構成比",
    ], height=46)

    total_sales = max(1, sum(to_int(p.get("subtotal_after_discount", 0)) for p in products))
    for i, p in enumerate(products, start=3):
        qty = to_int(p.get("quantity", 0))
        avg_price = to_int(p.get("avg_unit_price", 0))
        coupon_yen = to_int(p.get("couponYen", 0))
        points_yen = to_int(p.get("pointsYen", 0))
        sales_after = to_int(p.get("subtotal_after_discount", 0))
        cost_unit = to_int(p.get("cost_unit", 0))

        ws.cell(row=i, column=1, value=str(p.get("product_name", "")))
        ws.cell(row=i, column=2, value=avg_price)
        ws.cell(row=i, column=3, value=qty)
        ws.cell(row=i, column=4, value="=B{0}*C{0}".format(i))
        ws.cell(row=i, column=5, value=coupon_yen)
        ws.cell(row=i, column=6, value=points_yen)
        ws.cell(row=i, column=7, value="=D{0}-E{0}-F{0}".format(i))
        ws.cell(row=i, column=8, value=cost_unit)
        ws.cell(row=i, column=9, value="=H{0}*C{0}".format(i))
        ws.cell(row=i, column=10, value="=B{0}-H{0}".format(i))
        ws.cell(row=i, column=11, value="=G{0}-I{0}".format(i))
        ws.cell(row=i, column=12, value="=IF(G{0}=0,0,K{0}/G{0})".format(i))
        ws.cell(row=i, column=13, value=sales_after / total_sales if total_sales else 0)
        style_range(ws, i, i, {2: "yen", 3: "count", 4: "yen", 5: "yen", 6: "yen", 7: "yen",
                               8: "yen", 9: "yen", 10: "yen", 11: "yen", 12: "pct", 13: "pct"},
                    center_cols=(3,))
        ws.cell(row=i, column=7).font = Font(name="メイリオ", size=10, bold=True, color=NAVY)
        profit_unit = avg_price - cost_unit
        ws.cell(row=i, column=10).font = (GOOD_FONT if profit_unit >= 0 else BAD_FONT)
        profit = sales_after - cost_unit * qty
        ws.cell(row=i, column=11).font = (GOOD_FONT if profit >= 0 else BAD_FONT)
        if cost_unit <= 0:
            ws.cell(row=i, column=8).font = BAD_FONT
            ws.cell(row=i, column=9).font = BAD_FONT

    total_row = 3 + len(products)
    data_end = max(2, total_row - 1)
    ws.cell(row=total_row, column=1, value="合計")
    total_row_style(ws.cell(row=total_row, column=1), LEFT)
    for col, fmt in [(2, "yen"), (3, "count"), (4, "yen"), (5, "yen"), (6, "yen"), (7, "yen"),
                     (8, "yen"), (9, "yen"), (10, "yen"), (11, "yen"), (12, "pct"), (13, "pct")]:
        total_row_style(ws.cell(row=total_row, column=col), RIGHT, fmt)
    ws.row_dimensions[total_row].height = 24
    ws.cell(row=total_row, column=2, value="=IF(C{0}=0,0,D{0}/C{0})".format(total_row))
    ws.cell(row=total_row, column=3, value="=SUM(C3:C{0})".format(data_end))
    ws.cell(row=total_row, column=4, value="=SUM(D3:D{0})".format(data_end))
    ws.cell(row=total_row, column=5, value="=SUM(E3:E{0})".format(data_end))
    ws.cell(row=total_row, column=6, value="=SUM(F3:F{0})".format(data_end))
    ws.cell(row=total_row, column=7, value="=SUM(G3:G{0})".format(data_end))
    ws.cell(row=total_row, column=8, value="=IF(C{0}=0,0,I{0}/C{0})".format(total_row))
    ws.cell(row=total_row, column=9, value="=SUM(I3:I{0})".format(data_end))
    ws.cell(row=total_row, column=10, value="=B{0}-H{0}".format(total_row))
    ws.cell(row=total_row, column=11, value="=G{0}-I{0}".format(total_row))
    ws.cell(row=total_row, column=12, value="=IF(G{0}=0,0,K{0}/G{0})".format(total_row))
    ws.cell(row=total_row, column=13, value="=1")

    write_notes(ws, total_row + 1, DATA_COLS, [
        "※ 入金売上(割引後) ＝ 売上計算(割引前) − クーポン割引 − ポイント充当 です。",
        "※ 粗利 ＝ 入金売上(割引後) − 売上原価 です。売上原価 ＝ 仕入れ原価(単価) × 数量。",
        "※ 仕入れ原価が未登録（赤文字）の商品は原価0円で計算されるため、粗利が実際より大きく出ます。",
    ])

    if products:
        sales_chart = BarChart()
        sales_chart.type = "bar"
        sales_chart.style = 10
        sales_chart.title = "商品別入金売上"
        sales_chart.legend = None
        sales_chart.gapWidth = 55
        sales_chart.add_data(Reference(ws, min_col=7, min_row=2, max_row=total_row - 1), titles_from_data=True)
        sales_chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=total_row - 1))
        sales_chart.width = 11.8
        sales_chart.height = 8.0
        if sales_chart.series:
            sales_chart.series[0].graphicalProperties.solidFill = ACCENT
        ws.add_chart(sales_chart, "P3")

        profit_chart = BarChart()
        profit_chart.type = "bar"
        profit_chart.style = 12
        profit_chart.title = "商品別粗利"
        profit_chart.legend = None
        profit_chart.gapWidth = 55
        profit_chart.add_data(Reference(ws, min_col=11, min_row=2, max_row=total_row - 1), titles_from_data=True)
        profit_chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=total_row - 1))
        profit_chart.width = 11.8
        profit_chart.height = 8.0
        if profit_chart.series:
            profit_chart.series[0].graphicalProperties.solidFill = POS
        ws.add_chart(profit_chart, "V3")

        qty_chart = BarChart()
        qty_chart.type = "col"
        qty_chart.style = 11
        qty_chart.title = "商品別販売個数"
        qty_chart.legend = None
        qty_chart.gapWidth = 45
        qty_chart.add_data(Reference(ws, min_col=3, min_row=2, max_row=total_row - 1), titles_from_data=True)
        qty_chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=total_row - 1))
        qty_chart.width = 11.8
        qty_chart.height = 8.0
        if qty_chart.series:
            qty_chart.series[0].graphicalProperties.solidFill = "0EA5E9"
        ws.add_chart(qty_chart, "P22")

        pie = PieChart()
        pie.title = "売上構成比"
        pie.legend.position = "r"
        pie.varyColors = True
        pie.add_data(Reference(ws, min_col=7, min_row=3, max_row=total_row - 1), titles_from_data=False)
        pie.set_categories(Reference(ws, min_col=1, min_row=3, max_row=total_row - 1))
        if len(products) <= 6:
            pie.dataLabels = DataLabelList()
            pie.dataLabels.showPercent = True
            pie.dataLabels.showLeaderLines = True
        pie.width = 11.0
        pie.height = 8.8
        ws.add_chart(pie, "V22")

    ws.auto_filter.ref = "A2:M{0}".format(total_row)
    ws.freeze_panes = "A3"
    return ws


# ---------------------------------------------------------------------------
# 5. 注文一覧
# ---------------------------------------------------------------------------
def build_order_sheet(wb, orders):
    ws = wb.create_sheet("注文一覧")
    set_widths(ws, [30, 16, 15, 26, 12, 38, 12, 14, 15, 15, 12, 12, 11])
    write_title_bar(ws, 1, "注文一覧（支払いが完了した注文）", 13, height=34)
    write_headers(ws, 2, ["注文ID", "注文日時", "購入者名", "メール", "支払方法", "商品内訳", "小計(円)", "クーポンコード", "クーポン割引(円)", "ポイント充当(円)", "入金額(円)", "売上原価(円)", "粗利(円)"], height=44)

    for i, o in enumerate(orders, start=3):
        ws.cell(row=i, column=1, value=str(o.get("id", "")))
        ws.cell(row=i, column=2, value=str(o.get("created_at", "") or ""))
        ws.cell(row=i, column=3, value=str(o.get("name", "") or ""))
        ws.cell(row=i, column=4, value=str(o.get("email", "") or ""))
        ws.cell(row=i, column=5, value=str(o.get("payment_method", "") or ""))
        ws.cell(row=i, column=6, value=str(o.get("itemsText", "") or ""))
        ws.cell(row=i, column=7, value=to_int(o.get("subtotal", 0)))
        ws.cell(row=i, column=8, value=str(o.get("couponCode", "") or ""))
        ws.cell(row=i, column=9, value=to_int(o.get("coupon", 0)))
        ws.cell(row=i, column=10, value=to_int(o.get("points", 0)))
        ws.cell(row=i, column=11, value=to_int(o.get("total", 0)))
        ws.cell(row=i, column=12, value=to_int(o.get("cost", 0)))
        ws.cell(row=i, column=13, value="=K{0}-L{0}".format(i))

    n = len(orders)
    if n > 0:
        style_range(ws, 3, 2 + n, {7: "yen", 9: "yen", 10: "yen", 11: "yen", 12: "yen", 13: "yen"}, center_cols=(5,))
    total_row = 3 + n
    ws.cell(row=total_row, column=1, value="合計")
    total_row_style(ws.cell(row=total_row, column=1), LEFT)
    for col in (7, 9, 10, 11, 12, 13):
        cell = ws.cell(row=total_row, column=col, value="=SUM({0}3:{0}{1})".format(get_column_letter(col), max(2, total_row - 1)))
        total_row_style(cell, RIGHT, "yen")
    ws.row_dimensions[total_row].height = 24
    ws.freeze_panes = "A3"
    return ws


# ---------------------------------------------------------------------------
# 6. 処分履歴（廃棄ロス）
# ---------------------------------------------------------------------------
def build_disposal_sheet(wb, rows):
    ws = wb.create_sheet("処分履歴")
    headers = ["日時", "商品名", "商品ID", "数量", "処分理由", "メモ", "ロット", "原価(1個)", "処分原価(円)"]
    set_widths(ws, [18, 28, 9, 8, 16, 26, 16, 12, 14])
    write_title_bar(ws, 1, "処分履歴（廃棄ロス）", len(headers), height=34)
    write_headers(ws, 2, headers, height=30)
    for i, r in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=str(r.get("created_at") or ""))
        ws.cell(row=i, column=2, value=str(r.get("product_name") or ""))
        ws.cell(row=i, column=3, value=to_int(r.get("product_id", 0)))
        ws.cell(row=i, column=4, value=to_int(r.get("quantity", 0)))
        ws.cell(row=i, column=5, value=str(r.get("reason") or ""))
        ws.cell(row=i, column=6, value=str(r.get("memo") or ""))
        ws.cell(row=i, column=7, value=str(r.get("lot_label") or ""))
        ws.cell(row=i, column=8, value=to_int(r.get("cost", 0)))
        ws.cell(row=i, column=9, value=to_int(r.get("cost_total", 0)))

    data_end = 2 + len(rows)
    style_range(ws, 3, data_end, {4: "count", 8: "yen", 9: "yen"})
    total = sum(to_int(r.get("cost_total", 0)) for r in rows)
    total_row = 3 + len(rows)
    ws.cell(row=total_row, column=1, value="合計")
    total_row_style(ws.cell(row=total_row, column=1), LEFT)
    for col, fmt in [(4, "count"), (9, "yen")]:
        total_row_style(ws.cell(row=total_row, column=col), RIGHT, fmt)
    ws.cell(row=total_row, column=4, value="=SUM(D3:D{0})".format(max(2, data_end)))
    ws.cell(row=total_row, column=9, value="=SUM(I3:I{0})".format(max(2, data_end)))
    ws.row_dimensions[total_row].height = 24

    write_notes(ws, total_row + 2, len(headers), [
        "商品管理で「在庫を減らす」を行った記録です。処分した分の原価は「売上サマリー」の「廃棄ロス」に加算され、利益から差し引かれます。",
        "処分原価 ＝ 原価(1個) × 数量 です。",
        "この期間の廃棄ロス合計：{0} 円".format(total),
    ])
    ws.freeze_panes = "A3"
    return ws


def build_workbook(d):
    range_info = d.get("range") or {}
    summary = d.get("summary") or {}
    prev = d.get("prev")
    products = d.get("products") or []
    orders = d.get("orders") or []
    range_label = str(range_info.get("label") or "")
    now_jst = str(range_info.get("nowJst") or "")

    wb = Workbook()
    wb.remove(wb.active)
    build_overview_sheet(wb, range_label, now_jst, summary, products, prev)
    build_summary_sheet(wb, range_label, now_jst, summary)
    build_compare_sheet(wb, range_label, prev)
    build_product_sheet(wb, products)
    build_order_sheet(wb, orders)
    disposals = d.get("disposals") or []
    if disposals:
        build_disposal_sheet(wb, disposals)
    return wb


def build_xlsx_bytes(payload):
    wb = build_workbook(payload or {})
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    label = str((payload.get("range") or {}).get("label") or "期間") \
        .replace("\\", "_").replace("/", "_").replace(":", "_") \
        .replace("*", "_").replace("?", "_").replace('"', "_") \
        .replace("<", "_").replace(">", "_").replace("|", "_") \
        .replace("～", "_").replace("~", "_").strip() or "期間"
    fname = f"OFFICE NAGAZON売上_{label}.xlsx"
    return raw, fname


class handler(BaseHTTPRequestHandler):
    def do_OPTIONS(self):
        self.send_response(204)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self):
        self.send_response(405)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.end_headers()
        self.wfile.write(json.dumps({"error": "method_not_allowed"}).encode("utf-8"))

    def do_POST(self):
        try:
            length = int(self.headers.get("content-length", "0") or "0")
            raw_body = self.rfile.read(length) if length > 0 else b"{}"
            payload = json.loads(raw_body.decode("utf-8")) if raw_body else {}
            raw, fname = build_xlsx_bytes(payload)
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", "attachment; filename*=UTF-8''{}".format(quote(fname)))
            self.send_header("Content-Length", str(len(raw)))
            self.end_headers()
            self.wfile.write(raw)
        except Exception as e:
            self.send_response(500)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.end_headers()
            self.wfile.write(json.dumps({"error": str(e)}, ensure_ascii=False).encode("utf-8"))
